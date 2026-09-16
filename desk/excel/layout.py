"""Sheet scaffolding: titles, the SIM banner, and a `Table` that knows where each of its columns lives.

Every formula in this workbook is built from `Table.cell(column, i)` / `Table.abs(column)` rather than from
hand-counted letters, so inserting a column can never silently point a formula at the wrong data. A `Table` also
records the colour of each cell it writes (input / formula / pasted), which is what `desk.excel.build` counts to
publish the formula-coverage statistics in `docs/35_excel_workbook.md`.
"""

from __future__ import annotations

import datetime as dt
import math
from dataclasses import dataclass, field
from typing import Any, Iterable, Sequence

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from desk import SIM_LABEL
from desk.excel import style

INPUT, FORMULA, PASTED, KEY = "input", "formula", "pasted", "key"

_FONT = {INPUT: style.INPUT_FONT, FORMULA: style.FORMULA_FONT, PASTED: style.PASTED_FONT, KEY: style.KEY_FONT}


@dataclass
class Counts:
    """Cells written per kind, per sheet — the source of the coverage table in the methods doc."""

    per_sheet: dict[str, dict[str, int]] = field(default_factory=dict)

    def add(self, sheet: str, kind: str, n: int = 1) -> None:
        self.per_sheet.setdefault(sheet, {INPUT: 0, FORMULA: 0, PASTED: 0, KEY: 0})[kind] += n

    def ratio(self, sheet: str) -> float:
        c = self.per_sheet.get(sheet, {})
        total = sum(c.values())
        return c.get(FORMULA, 0) / total if total else 0.0


@dataclass
class Column:
    name: str
    kind: str = FORMULA
    fmt: str | None = None
    width: float | None = None
    header: str | None = None


class Table:
    """A header row plus `n_rows` data rows, addressable by column name."""

    def __init__(self, sheet: "Sheet", columns: Sequence[Column], header_row: int, first_col: int = 1) -> None:
        self.sheet = sheet
        self.columns = list(columns)
        self.header_row = header_row
        self.first_row = header_row + 1
        self.first_col = first_col
        self.n_rows = 0
        self._index = {c.name: first_col + i for i, c in enumerate(self.columns)}
        ws = sheet.ws
        for c in self.columns:
            j = self._index[c.name]
            cell = ws.cell(header_row, j, c.header or c.name)
            cell.font = style.HEADER_FONT
            cell.fill = style.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = style.BOX
            ws.column_dimensions[get_column_letter(j)].width = c.width or _auto_width(c.header or c.name)
        ws.row_dimensions[header_row].height = 30
        sheet.counts.add(sheet.name, KEY, len(self.columns))

    # ------------------------------------------------------------------------------------------ addressing
    def col(self, name: str) -> str:
        return get_column_letter(self._index[name])

    def col_index(self, name: str) -> int:
        return self._index[name]

    def cell(self, name: str, i: int) -> str:
        """A1 reference of data row `i` (0-based) — relative, for use inside the same row."""
        return f"{self.col(name)}{self.first_row + i}"

    def rowref(self, name: str, i: int) -> str:
        return f"${self.col(name)}{self.first_row + i}"

    def abs(self, name: str, n_rows: int | None = None) -> str:
        n = self.n_rows if n_rows is None else n_rows
        c = self.col(name)
        return f"{self.sheet.name}!${c}${self.first_row}:${c}${self.first_row + max(n, 1) - 1}"

    def abs_block(self, first: str, last: str, n_rows: int | None = None) -> str:
        n = self.n_rows if n_rows is None else n_rows
        return (f"{self.sheet.name}!${self.col(first)}${self.first_row}:"
                f"${self.col(last)}${self.first_row + max(n, 1) - 1}")

    def header_abs(self, first: str, last: str) -> str:
        return (f"{self.sheet.name}!${self.col(first)}${self.header_row}:"
                f"${self.col(last)}${self.header_row}")

    # ---------------------------------------------------------------------------------------------- writing
    def write_row(self, i: int, values: dict[str, Any], kinds: dict[str, str] | None = None) -> None:
        ws = self.sheet.ws
        kinds = kinds or {}
        r = self.first_row + i
        for c in self.columns:
            v0 = values.get(c.name)
            if _blank(v0):
                continue     # an empty cell is neither a formula nor an input, and openpyxl does not persist one
            v = values[c.name]
            kind = kinds.get(c.name, c.kind)
            # A leading "=" means a formula only in a column declared as one. Register notes and source strings can
            # legitimately begin with "=" ("= ₹250 per 5 MT lot"), and Excel would silently try to compile them.
            is_formula = isinstance(v, str) and v.startswith("=") and kind == FORMULA
            if kind == FORMULA and not is_formula:
                kind = KEY          # a label written into a computed column is a label, and is counted as one
            cell = ws.cell(r, self._index[c.name], _clean(v))
            if isinstance(v, str) and not is_formula:
                cell.data_type = "s"
            cell.font = _FONT[kind]
            fmt = c.fmt or style.fmt_for(c.name)
            if isinstance(v, (dt.date, dt.datetime)):
                fmt = style.FMT["date"]
            elif isinstance(v, bool) or (isinstance(v, str) and not v.startswith("=")):
                fmt = style.FMT["text"] if not isinstance(v, bool) else "General"
            cell.number_format = fmt
            self.sheet.counts.add(self.sheet.name, kind)
        self.n_rows = max(self.n_rows, i + 1)

    def write_rows(self, rows: Iterable[dict[str, Any]], kinds: dict[str, str] | None = None) -> None:
        for i, r in enumerate(rows):
            self.write_row(i, r, kinds)

    def freeze(self, at_col: str | None = None) -> None:
        col = self.col(at_col) if at_col else "A"
        self.sheet.ws.freeze_panes = f"{col}{self.first_row}"

    def autofilter(self) -> None:
        last = get_column_letter(self.first_col + len(self.columns) - 1)
        first = get_column_letter(self.first_col)
        self.sheet.ws.auto_filter.ref = (f"{first}{self.header_row}:{last}"
                                         f"{self.first_row + max(self.n_rows, 1) - 1}")


class Sheet:
    """One worksheet: SIM banner, title, and a cursor that section/table writers advance."""

    def __init__(self, wb, name: str, title: str, counts: Counts, subtitle: str = "") -> None:
        self.wb = wb
        self.name = name
        self.ws: Worksheet = wb.create_sheet(name)
        self.counts = counts
        self.counts.add(name, KEY, 0)
        ws = self.ws
        ws.sheet_view.showGridLines = False
        b = ws.cell(1, 1, SIM_LABEL)
        b.font = style.SIM_FONT
        b.fill = style.SIM_FILL
        t = ws.cell(2, 1, title)
        t.font = style.TITLE_FONT
        self.row = 3
        self.counts.add(name, KEY, 2)
        if subtitle:
            s = ws.cell(3, 1, subtitle)
            s.font = style.NOTE_FONT
            self.row = 4
            self.counts.add(name, KEY)

    def note(self, text: str, row: int | None = None) -> int:
        r = self.row if row is None else row
        c = self.ws.cell(r, 1, text)
        c.font = style.NOTE_FONT
        self.counts.add(self.name, KEY)
        self.row = r + 1
        return r

    def section(self, text: str) -> int:
        r = self.row + 1
        self.section_at(r, text)
        self.row = r + 1
        return r

    def section_at(self, row: int, text: str) -> int:
        c = self.ws.cell(row, 1, text)
        c.font = style.SECTION_FONT
        self.counts.add(self.name, KEY)
        return row

    def put(self, row: int, col: int, value, kind: str = KEY, fmt: str | None = None):
        c = self.ws.cell(row, col, value)
        c.font = _FONT[kind]
        if fmt:
            c.number_format = fmt
        self.counts.add(self.name, kind)
        return c

    def table(self, columns: Sequence[Column], first_col: int = 1, gap: int = 1) -> Table:
        header = self.row + gap
        t = Table(self, columns, header, first_col)
        self.row = t.first_row
        return t

    def after(self, table: Table, gap: int = 2) -> None:
        self.row = table.first_row + max(table.n_rows, 1) - 1 + gap


def _auto_width(header: str) -> float:
    return min(26.0, max(9.0, len(header) * 1.05 + 2))


def _blank(v: Any) -> bool:
    """True for anything openpyxl would not persist: None, the empty string, NaN/inf, a pandas NaT."""
    if v is None:
        return True
    if isinstance(v, str):
        return not v
    if isinstance(v, bool):
        return False
    if isinstance(v, float):
        return math.isnan(v) or math.isinf(v)
    return type(v).__name__ == "NaTType"


def _clean(v: Any) -> Any:
    return None if _blank(v) else v
