"""Recalculate the saved workbook from its formulas alone and prove it still agrees with the Python engine.

`desk.excel.build` writes formulas without cached values, so nothing here can be reading a number openpyxl stored:
the `formulas` library parses every formula, builds the dependency graph and evaluates it. What the tests assert:

1. **Every `check` cell reads PASS** — parity row by row, both Table 1.5 grids cell by cell, the term structure, the
   cash ledger, the MCX margin schedule, every leg mark on the marking grid, per-trade cumulative P&L at
   `WINDOW_END` and `HORIZON_END`, the book equity curve, and the attribution identities — and so does the
   scoreboard formula that aggregates them.
2. **The computed sheets really are computed** — a formula-vs-constant census per sheet, with a floor per sheet, so
   a future edit cannot quietly replace a formula with a pasted number.
3. **No cell recalculates to an error** — `#REF!`, `#NAME?`, `#DIV/0!`, `#VALUE!`, `#N/A`, `#NULL!`, `#NUM!`.
4. **The workbook is byte-for-byte deterministic** (CONTRACTS §1.5): rebuilding it produces an identical file.

Scope: the **whole workbook** is recalculated — 127k cells, about two minutes — so nothing is sampled and no result
is quoted from a subset. The run writes `outputs/excel/reconciliation.json`, including the SHA-256 and byte size of
the exact workbook it recalculated (hashed before the recalculation and re-checked before writing). Everything that
quotes the result — `docs/35_excel_workbook.md` §7, the interview pack, the README checklist — goes through
`desk.excel.reconciliation_status`, which reports VERIFIED only while that hash matches the workbook on disk (else
STALE, NOT_RUN or FAILED); re-run `desk.excel.build` and `run_all.py --only P6` after the test to refresh them.
"""

from __future__ import annotations

import gc
import hashlib
import json
import re
import time

import pytest

from desk.excel import build, sources
from desk.excel import reconciliation_status as rs

WORKBOOK = build.WORKBOOK
RECON_JSON = rs.RECON_JSON

ERROR_VALUES = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NULL!", "#NUM!")

# Minimum share of written cells that must be formulas, per computed sheet. Registers (Inputs, Counterparties,
# Trade_Book) are deliberately input-heavy and are checked by an absolute formula count instead.
FORMULA_FLOOR = {
    "Market_Daily": 0.25,
    "Market_Weekly": 0.75,
    "Parity": 0.85,
    "Sensitivity": 0.70,
    "Term_Structure": 0.50,
    "Cashflows": 0.30,
    "MTM_Daily": 0.20,
    "Attribution": 0.22,
    "Equity_Curve": 0.40,
    "Checks": 0.45,
}
MIN_FORMULAS = {
    "Market_Daily": 2000, "Market_Weekly": 700, "Parity": 12000, "Sensitivity": 6000, "Term_Structure": 600,
    "Trade_Book": 40, "Cashflows": 10000, "MTM_Daily": 6000, "Attribution": 800, "Equity_Curve": 800,
    "Checks": 3000,
}
CELL_KEY = re.compile(r"\]([A-Za-z_0-9]+)'!([A-Z]+\d+)$")


# --------------------------------------------------------------------------------------------------- fixtures
@pytest.fixture(scope="module")
def workbook():
    import openpyxl

    assert WORKBOOK.exists(), (f"{WORKBOOK} missing — run "
                               f"DESK_OFFLINE=1 .venv/bin/python -c \"import desk.excel.build as m; m.main()\"")
    return openpyxl.load_workbook(WORKBOOK)


@pytest.fixture(scope="module")
def census(workbook):
    """formula vs constant cells per sheet, straight off the saved file."""
    out = {}
    for ws in workbook:
        f = c = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if cell.data_type == "f":
                    f += 1
                else:
                    c += 1
        out[ws.title] = {"formula": f, "constant": c, "ratio": f / (f + c) if (f + c) else 0.0}
    return out


@pytest.fixture(scope="module")
def recalc():
    """Full-workbook recalculation: {SHEET_UPPER: {coord: value}} plus timing.

    The cycle collector is switched off for the duration. `formulas` builds a dependency graph of millions of small
    objects, almost all of which survive every generation, so the automatic collections cost several times the
    calculation itself — the same run takes ~2 minutes with the collector off and ~15 with it on.
    """
    import formulas

    fingerprint = rs.workbook_record(WORKBOOK)         # what is about to be recalculated, before it is read
    gc.disable()
    try:
        t0 = time.time()
        xl = formulas.ExcelModel().loads(str(WORKBOOK)).finish()
        solution = xl.calculate()
    finally:
        gc.enable()
    values: dict[str, dict[str, object]] = {}
    for key, node in solution.items():
        m = CELL_KEY.search(key)
        if not m:
            continue
        try:
            v = node.value[0, 0]
        except Exception:                                    # ranges and non-cell nodes
            continue
        if hasattr(v, "item"):
            try:
                v = v.item()
            except Exception:
                v = str(v)
        if not isinstance(v, (int, float, str, bool)) and v is not None:
            v = str(v)
        values.setdefault(m.group(1), {})[m.group(2)] = v
    return {"values": values, "seconds": time.time() - t0,
            "n_cells": sum(len(v) for v in values.values()), "workbook": fingerprint}


def _cell(workbook, recalc, sheet: str, coord: str):
    """The recalculated value of a formula cell, or the stored value of a constant."""
    c = workbook[sheet][coord]
    if c.data_type == "f":
        return recalc["values"].get(sheet.upper(), {}).get(coord)
    return c.value


def _header_row(ws, name: str) -> int:
    for r in range(1, min(ws.max_row, 60) + 1):
        for col in range(1, ws.max_column + 1):
            if ws.cell(r, col).value == name:
                return r
    raise AssertionError(f"header {name!r} not found on {ws.title}")


# ------------------------------------------------------------------------------------------------- the tests
@pytest.mark.slow
def test_recalculation_has_no_error_values(recalc):
    bad = [(sheet, coord, v) for sheet, cells in recalc["values"].items() for coord, v in cells.items()
           if isinstance(v, str) and v in ERROR_VALUES]
    assert not bad, f"{len(bad)} error cells after recalculation, e.g. {bad[:10]}"


@pytest.mark.slow
def test_every_check_cell_passes(workbook, recalc):
    fails = []
    n = 0
    for sheet, cells in recalc["values"].items():
        for coord, v in cells.items():
            if v in ("PASS", "FAIL"):
                n += 1
                if v == "FAIL":
                    fails.append(f"{sheet}!{coord}")
    assert n > 3000, f"only {n} PASS/FAIL cells recalculated — the check columns are missing"
    assert not fails, f"{len(fails)} failing check cells, e.g. {fails[:20]}"


@pytest.mark.slow
def test_scoreboard_recalculates_to_pass(workbook, recalc):
    ws = workbook["Checks"]
    hr = _header_row(ws, "family")
    families = {}
    r = hr + 1
    while ws.cell(r, 1).value and ws.cell(r, 1).value != "ALL CHECKS":
        families[str(ws.cell(r, 1).value)] = {
            "n_rows": ws.cell(r, 2).value,
            "tolerance": ws.cell(r, 3).value,
            "max_abs_diff": _cell(workbook, recalc, "Checks", f"D{r}"),
            "n_fail": _cell(workbook, recalc, "Checks", f"E{r}"),
            "status": _cell(workbook, recalc, "Checks", f"F{r}"),
        }
        r += 1
    assert families, "no scoreboard rows found on Checks"
    assert all(f["status"] == "PASS" for f in families.values()), families
    verdict = _cell(workbook, recalc, "Checks", f"B{r}")
    assert verdict == "PASS", f"ALL CHECKS = {verdict!r}: {families}"
    _write_report(workbook, recalc, families)


def test_published_reconciliation_report_passes():
    """The workbook's published control — the recalculated scoreboard — as the slow run last wrote it, tied by SHA-256
    to the workbook on disk.

    Fast by design so the default suite covers it. It fails when the record passed for a *different* workbook (STALE)
    or did not pass (FAILED), so a rebuilt workbook cannot keep quoting an old verification; it skips only when no
    record exists (NOT_RUN), which every consumer then states in so many words.
    """
    st = rs.status()
    if st.status == rs.NOT_RUN:
        pytest.skip(f"{st.reason} {st.how_to_verify()}")
    assert st.status == rs.VERIFIED, f"{st.status}: {st.reason} {st.how_to_verify()}"
    rep_ = st.report
    assert rep_["scoreboard"] == "PASS"
    assert rep_["n_check_failures"] == 0 and rep_["n_errors"] == 0
    assert rep_["n_check_cells"] > 3000 and rep_["scope"] == "full workbook"
    assert rep_["families"] and all(f["status"] == "PASS" for f in rep_["families"].values()), rep_["families"]
    assert rep_["workbook"]["path"] == "outputs/excel/Metals_Desk_Master.xlsx"


def test_report_writer_records_the_hash_of_the_recalculated_workbook(tmp_path, monkeypatch):
    """`_write_report` (called by the slow scoreboard test) stamps the workbook's SHA-256 and size into the record, and
    refuses to write if the file changed during the recalculation. Exercised on a tiny workbook, no recalculation."""
    import openpyxl

    wb_path, recon_path = tmp_path / "tiny.xlsx", tmp_path / "reconciliation.json"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.active["A2"] = "=A1*2"
    wb.save(wb_path)
    monkeypatch.setitem(globals(), "WORKBOOK", wb_path)
    monkeypatch.setitem(globals(), "RECON_JSON", recon_path)
    recalc = {"values": {"SHEET": {"A2": 2, "B1": "PASS"}}, "seconds": 1.0, "n_cells": 1,
              "workbook": rs.workbook_record(wb_path)}
    families = {"toy": {"max_abs_diff": 0.0, "tolerance": "tol", "status": "PASS"}}
    _write_report(openpyxl.load_workbook(wb_path), recalc, families)

    written = json.loads(recon_path.read_text(encoding="utf-8"))
    assert written["workbook"]["sha256"] == hashlib.sha256(wb_path.read_bytes()).hexdigest()
    assert written["workbook"]["bytes"] == wb_path.stat().st_size
    assert written["workbook"]["hash_source"] == rs.HASH_RECORDED_BY_RUN
    assert rs.status(recon_path, wb_path).status == rs.VERIFIED

    with wb_path.open("ab") as fh:                       # the workbook changes after it was recalculated
        fh.write(b"\0")
    assert rs.status(recon_path, wb_path).status == rs.STALE
    with pytest.raises(AssertionError, match="changed on disk"):
        _write_report(wb, recalc, families)


def test_status_helper_points_at_this_workbook_and_this_test():
    assert rs.WORKBOOK == build.WORKBOOK and rs.RECON_JSON == RECON_JSON
    assert "DESK_RUN_SLOW=1" in rs.VERIFY_COMMAND and "tests/test_excel_reconciliation.py" in rs.VERIFY_COMMAND


def test_computed_sheets_are_formula_driven(census):
    """A future edit cannot quietly swap a formula for a pasted number."""
    for sheet, floor in FORMULA_FLOOR.items():
        got = census[sheet]["ratio"]
        assert got >= floor, f"{sheet}: formula share {got:.0%} below the {floor:.0%} floor"
    for sheet, minimum in MIN_FORMULAS.items():
        assert census[sheet]["formula"] >= minimum, f"{sheet}: only {census[sheet]['formula']} formula cells"
    assert census["Parity"]["formula"] > 0 and census["Sensitivity"]["formula"] > 0


def test_no_cached_values_are_relied_on():
    """openpyxl writes formulas without cached results, so every number the tests below read is recalculated.

    A cached result would appear in the sheet XML as a non-empty `<v>` element immediately after the `</f>` that
    produced it.
    """
    import zipfile

    with zipfile.ZipFile(WORKBOOK) as z:
        sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
        assert sheets
        with_formulas = 0
        for name in sheets:
            xml = z.read(name).decode("utf-8", "replace")
            # Evaluate to a bool BEFORE asserting. `assert "</f><v>" not in xml` on a multi-megabyte sheet made pytest's
            # assertion rewriting run difflib over the whole XML when it failed — over ten minutes of CPU for one
            # failing check, which is most of why the suite had become so slow.
            # openpyxl writes an EMPTY `<v></v>` after some formulas; that carries no cached result, so only a
            # non-empty value counts (the old substring test failed on the empty element).
            cached = len(re.findall(r"</f><v>[^<]+</v>", xml))
            assert cached == 0, f"{name} carries {cached} cached formula results"
            with_formulas += "<f>" in xml
    assert with_formulas >= 10, f"only {with_formulas} sheets contain formulas"


def test_named_ranges_are_used_and_resolve(workbook):
    """Key parameters are addressed by name. (A name that did not resolve would recalculate to #NAME?, which the slow
    `test_recalculation_has_no_error_values` catches; this check needs no recalculation.)"""
    names = set(workbook.defined_names)
    for expected in ("bcd_scrap_hs7602", "sws_rate_on_bcd", "igst_rate_hs7602", "conversion_cost_inr_t",
                     "margin_threshold_inr_t", "conv_step_5a", "param_keys", "param_values",
                     "mcx_al_margin_used_frac", "tol_parity_inr_t"):
        assert expected in names, f"named range {expected} missing"
    parity = workbook["Parity"]
    hr = _header_row(parity, "net_arb_inr_t")
    formula = parity.cell(hr + 1, _col(parity, hr, "bcd_inr_t")).value
    assert "bcd_scrap_hs7602" in formula, formula


def _col(ws, header_row: int, name: str) -> int:
    for c in range(1, ws.max_column + 1):
        if ws.cell(header_row, c).value == name:
            return c
    raise AssertionError(f"column {name!r} not found")


# ------------------------------------------------------------------------------------- report for the methods doc
def _write_report(workbook, recalc, families) -> None:
    # The record names the exact bytes it verified; refuse to write it if the file changed while it was recalculated.
    now = rs.fingerprint(WORKBOOK)
    assert now["sha256"] == recalc["workbook"]["sha256"], (
        f"{WORKBOOK} changed on disk during the recalculation — re-run the test with nothing else writing the workbook")
    census = {}
    for ws in workbook:
        f = c = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if cell.data_type == "f":
                    f += 1
                else:
                    c += 1
        if f:
            census[ws.title] = {"formula": f, "constant": c, "ratio": f / (f + c)}
    n_check = sum(1 for cells in recalc["values"].values() for v in cells.values() if v in ("PASS", "FAIL"))
    n_fail = sum(1 for cells in recalc["values"].values() for v in cells.values() if v == "FAIL")
    n_err = sum(1 for cells in recalc["values"].values() for v in cells.values()
                if isinstance(v, str) and v in ERROR_VALUES)
    report = {
        "scope": "full workbook",
        "scope_note": "Every formula cell in every sheet is recalculated; nothing is sampled.",
        "n_recalculated": recalc["n_cells"],
        "seconds": round(recalc["seconds"], 1),
        "n_check_cells": n_check,
        "n_check_failures": n_fail,
        "n_errors": n_err,
        "scoreboard": "PASS" if n_fail == 0 else "FAIL",
        "sheets": census,
        "workbook": recalc["workbook"],
        "families": {k: {"max_abs_diff": float(v["max_abs_diff"]), "tolerance": str(v["tolerance"]),
                         "status": str(v["status"])} for k, v in families.items()},
    }
    RECON_JSON.write_text(json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(f"\nrecalculated {report['n_recalculated']:,} cells in {report['seconds']:.0f}s — "
          f"{n_check:,} check cells, {n_fail} failures, {n_err} error values — workbook SHA-256 "
          f"{report['workbook']['sha256']}")
    for sheet, c in sorted(census.items()):
        print(f"  {sheet:<16} formula {c['formula']:>6,}  constant {c['constant']:>6,}  ratio {c['ratio']:.0%}")
    for name, f in families.items():
        print(f"  {f['status']:<5} max|diff| {float(f['max_abs_diff']):>14,.4f}  ({f['tolerance']})  {name}")


# ------------------------------------------------------------------------------ determinism (runs last: it rebuilds)
@pytest.mark.slow
def test_workbook_is_byte_for_byte_deterministic(tmp_path):
    """CONTRACTS §1.5: re-running the stage reproduces the file exactly (fixed zip stamps and doc properties).

    Runs last because it rebuilds the whole workbook in this process; the frames it loads would otherwise sit in
    memory beside the recalculation graph and slow the recalculation down several-fold.
    """
    again = tmp_path / "again.xlsx"
    wb, _counts, _ctx = build.build()
    build.save(wb, again)
    got = hashlib.sha256(again.read_bytes()).hexdigest()
    want = hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()
    sources.load.cache_clear()
    gc.collect()
    assert got == want, ("rebuilding the workbook produced a different file — something in desk.excel is not "
                         "deterministic (a clock, a set iteration, an unsorted dict)")
